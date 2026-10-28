"""The valuation engine's market inputs (Spec 02A §5.7; job 2A-05): parsers for Damodaran's data sets and the IMF
DataMapper, the WEO edition rule, and the two helpers the cost-of-capital job (2A-13) uses.

Sources (all read from raw files the task stored; nothing here fetches):
- Damodaran, `ERPbymonth.xlsx`, sheet "Historical ERP": the implied equity risk premium of the S&P 500, one row a month.
  The series is the column "ERP (T12 m with sustainable payout)" (Damodaran's headline implied premium: October 2026
  is 4.20%, the figure on his page; Data Audit 02 E4). Never the total-ERP column of `ctryprem` (that adds the country
  premium; §5.7 adds India's separately).
- Damodaran, `ctryprem*.xlsx`, sheet "ERPs by country": one edition per file, dated inside the file ("Date of update:"),
  not by its name; India's row gives the rating, the rating-based default spread and the country risk premium. The header
  row sits about 65 rows above the India row (Data Audit 02a finding 13), so rows are found by label.
- Damodaran, `betaIndia.xls` (BIFF, read with xlrd), sheet "Industry Averages": one row per industry, dated by "Date
  updated:" inside the file.
- IMF DataMapper API: `{"values": {<indicator>: {<country>: {<year>: <value>}}}, "api": {...}}`. The reply doesn't name
  the WEO edition, so the vintage comes from `config/equity/weo_editions.yaml` (see `choose_vintage`).

Damodaran's files are spreadsheets people edit: a date may be text ("1-Sep-24"), a percentage may be text ("4.05%"),
and a cell may hold a note. The parsers read those and refuse what they can't place, naming the row.

Percentages are stored as percent (4.20 for 4.20%), rounded to six decimals, as the series' unit says.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import yaml

CONFIG_DIR = Path(__file__).resolve().parents[3] / "config" / "equity"

DAMODARAN_BASE = "https://pages.stern.nyu.edu/~adamodar/"
DATACURRENT_URL = DAMODARAN_BASE + "New_Home_Page/datacurrent.html"
ERP_MONTHLY_URL = DAMODARAN_BASE + "pc/implprem/ERPbymonth.xlsx"
CTRYPREM_DEFAULT_URL = DAMODARAN_BASE + "pc/datasets/ctryprem.xlsx"
BETA_INDIA_URL = DAMODARAN_BASE + "pc/datasets/betaIndia.xls"
IMF_URL = "https://www.imf.org/external/datamapper/api/v1/{indicator}/{country}"

ERP_FROM = date(2019, 1, 1)                     # "every month the file holds from Jan 2019" (the card)
ERP_COLUMN = "ERP (T12 m with sustainable payout)"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS_TYPE = "application/vnd.ms-excel"

MONTHS = {m: i for i, m in enumerate(
    ["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


class InputsError(ValueError):
    """A file that isn't in the shape the parser expects. The message says what was missing."""


# ------------------------------------------------------------------------------------------------ small readers
def pct(fraction: float) -> Decimal:
    """A fraction as percent, six decimals: 0.0174988 -> 1.749880."""
    return Decimal(repr(round(fraction * 100, 6)))


def _number(v) -> float | None:
    """A cell as a number: a float, or text such as '4.05%' (as a fraction) or '0.0405'; None for anything else."""
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    if isinstance(v, str):
        s = v.strip().replace(",", ".")
        try:
            return float(s[:-1]) / 100 if s.endswith("%") else float(s)
        except ValueError:
            return None
    return None


def _date(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, str):
        for fmt in ("%d-%b-%y", "%d-%b-%Y", "%Y-%m-%d", "%B %d, %Y", "%d %B %Y"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _norm(s) -> str:
    return re.sub(r"\s+", " ", str(s or "")).strip().lower()


def _xlsx_rows(body: bytes, sheet: str) -> list[tuple]:
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True, read_only=True)
    except Exception as e:                       # not a workbook (an HTML error page, a truncated download)
        raise InputsError(f"not an .xlsx workbook ({type(e).__name__})") from e
    if sheet not in wb.sheetnames:
        raise InputsError(f"no sheet {sheet!r} (sheets: {', '.join(wb.sheetnames[:6])})")
    return list(wb[sheet].iter_rows(values_only=True))


# ------------------------------------------------------------------------------------------------ Damodaran: ERP
def parse_erp_by_month(body: bytes) -> list[tuple[date, Decimal]]:
    """Every month from Jan 2019 the file holds: (the month's first day, the implied ERP in percent). A month whose
    cell can't be read as a number is an error (the series has no gaps), as is a repeated month."""
    rows = _xlsx_rows(body, "Historical ERP")
    head = next((i for i, r in enumerate(rows) if r and _norm(r[0]) == "start of month"), None)
    if head is None:
        raise InputsError("Historical ERP: no 'Start of month' header row")
    col = next((j for j, c in enumerate(rows[head]) if _norm(c) == _norm(ERP_COLUMN)), None)
    if col is None:
        raise InputsError(f"Historical ERP: no column {ERP_COLUMN!r}")
    out: dict[date, Decimal] = {}
    for r in rows[head + 1:]:
        month = _date(r[0]) if r else None
        if month is None or month < ERP_FROM:
            continue
        month = month.replace(day=1)
        v = _number(r[col])
        if v is None:
            raise InputsError(f"Historical ERP: {month:%b %Y} has no readable value ({r[col]!r})")
        if month in out:
            raise InputsError(f"Historical ERP: {month:%b %Y} appears twice")
        out[month] = pct(v)
    if not out:
        raise InputsError("Historical ERP: no month from Jan 2019")
    return sorted(out.items())


def month_gaps(months: list[date]) -> list[date]:
    """The months missing between the first and last of `months`."""
    have = set(months)
    gaps, m = [], min(months)
    while m < max(months):
        if m not in have:
            gaps.append(m)
        m = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    return gaps


# ------------------------------------------------------------------------------------------------ Damodaran: ctryprem
@dataclass(frozen=True)
class CountryPremium:
    edition: date                  # "Date of update:" inside the file
    country: str
    rating: str | None             # Moody's
    default_spread: Decimal        # percent
    crp: Decimal                   # percent
    total_erp: Decimal | None      # percent (shown for the audit; the engine never uses it)
    mature_erp: Decimal | None     # percent: the mature-market premium the edition used


def parse_ctryprem(body: bytes, country: str = "India") -> CountryPremium:
    rows = _xlsx_rows(body, "ERPs by country")
    edition = mature = None
    head = None
    for i, r in enumerate(rows):
        if not r or r[0] is None:
            continue
        label = _norm(r[0])
        if label.startswith("date of update"):
            edition = _date(r[1])
        elif label == "enter the current risk premium for a mature equity market":
            mature = _number(next((c for c in r[1:] if _number(c) is not None), None))
        elif label == "country" and head is None:
            head = i
            break
    if edition is None:
        raise InputsError("ERPs by country: no 'Date of update:' row with a date")
    if head is None:
        raise InputsError("ERPs by country: no header row ('Country' ...)")
    cols = {_norm(c): j for j, c in enumerate(rows[head]) if c is not None}
    need = ("moody's rating", "rating-based default spread", "country risk premium", "total equity risk premium")
    missing = [n for n in need if n not in cols]
    if missing:
        raise InputsError(f"ERPs by country: no column {', '.join(missing)}")
    row = next((r for r in rows[head + 1:] if r and _norm(r[0]) == _norm(country)), None)
    if row is None:
        raise InputsError(f"ERPs by country: no row for {country}")
    spread, crp = _number(row[cols["rating-based default spread"]]), _number(row[cols["country risk premium"]])
    if spread is None or crp is None:
        raise InputsError(f"ERPs by country: {country}'s spread or premium isn't a number")
    rating = row[cols["moody's rating"]]
    total = _number(row[cols["total equity risk premium"]])
    return CountryPremium(edition, country, str(rating).strip() if rating else None, pct(spread), pct(crp),
                          pct(total) if total is not None else None, pct(mature) if mature is not None else None)


_HREF = re.compile(r"""href\s*=\s*["']([^"'#?]*ctryprem[^"'#?]*\.xlsx?)(?:[?#][^"']*)?["']""", re.IGNORECASE)
_BETA_HREF = re.compile(r"""href\s*=\s*["']([^"'#?]*betaIndia\.xlsx?)(?:[?#][^"']*)?["']""", re.IGNORECASE)


def _absolute(base: str, href: str) -> str:
    from urllib.parse import urljoin
    return urljoin(base, href)


def find_ctryprem_links(html: str, base: str = DATACURRENT_URL) -> list[str]:
    """Every country-risk-premium workbook the page links, as absolute URLs, in page order, without repeats."""
    seen: list[str] = []
    for href in _HREF.findall(html):
        url = _absolute(base, href)
        if url not in seen:
            seen.append(url)
    return seen


def find_beta_india_link(html: str, base: str = DATACURRENT_URL) -> str | None:
    m = _BETA_HREF.search(html)
    return _absolute(base, m.group(1)) if m else None


def edition_from_name(url_or_name: str) -> date | None:
    """The month a dated file name stands for: `ctrypremJuly26.xlsx` -> 1 Jul 2026 (None for `ctryprem.xlsx`).
    A cross-check on the date inside the file, never the date itself."""
    m = re.search(r"ctryprem([A-Za-z]+)(\d{2})\.xlsx?$", url_or_name.rsplit("/", 1)[-1], re.IGNORECASE)
    if not m:
        return None
    month = MONTHS.get(m.group(1)[:3].lower())
    return date(2000 + int(m.group(2)), month, 1) if month else None


# ------------------------------------------------------------------------------------------------ Damodaran: betaIndia
@dataclass(frozen=True)
class IndustryRow:
    industry: str
    n_firms: int | None
    beta: Decimal | None
    de_ratio: Decimal | None
    effective_tax: Decimal | None
    unlevered_beta: Decimal | None
    cash_firm_value: Decimal | None
    unlevered_beta_cash_corrected: Decimal | None


BETA_COLUMNS = {"number of firms": "n_firms", "beta": "beta", "d/e ratio": "de_ratio",
                "effective tax rate": "effective_tax", "unlevered beta": "unlevered_beta",
                "cash/firm value": "cash_firm_value",
                "unlevered beta corrected for cash": "unlevered_beta_cash_corrected"}


def _dec(v) -> Decimal | None:
    n = _number(v)
    return None if n is None else Decimal(repr(n))


def parse_beta_india(body: bytes) -> tuple[date, list[IndustryRow]]:
    """The edition date and every row of the Industry Averages sheet (industries and Damodaran's 'Total Market' rows).
    Fractions stay as filed (D/E 0.30 is 30%), the columns of Spec 02A §3.2."""
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=body)
    except Exception as e:
        raise InputsError(f"not an .xls workbook ({type(e).__name__})") from e
    if "Industry Averages" not in book.sheet_names():
        raise InputsError(f"no sheet 'Industry Averages' (sheets: {', '.join(book.sheet_names())})")
    sh = book.sheet_by_name("Industry Averages")
    rows = [sh.row_values(i) for i in range(sh.nrows)]
    edition = None
    for r in rows[:6]:
        if r and _norm(r[0]).startswith("date updated") and isinstance(r[1], float):
            edition = (datetime(1899, 12, 30) + timedelta(days=r[1])).date()
    if edition is None:
        raise InputsError("Industry Averages: no 'Date updated:' date")
    head = next((i for i, r in enumerate(rows) if r and _norm(r[0]) == "industry name"), None)
    if head is None:
        raise InputsError("Industry Averages: no 'Industry Name' header row")
    cols = {BETA_COLUMNS[_norm(c)]: j for j, c in enumerate(rows[head]) if _norm(c) in BETA_COLUMNS}
    if len(cols) != len(BETA_COLUMNS):
        raise InputsError("Industry Averages: missing columns " + ", ".join(sorted(set(BETA_COLUMNS.values()) - set(cols))))
    out, seen = [], set()
    for r in rows[head + 1:]:
        name = str(r[0]).strip() if r and r[0] not in ("", None) else ""
        if not name:
            continue
        if name in seen:
            raise InputsError(f"Industry Averages: {name!r} appears twice")
        seen.add(name)
        n = _number(r[cols["n_firms"]])
        out.append(IndustryRow(name, int(n) if n is not None else None,
                               *(_dec(r[cols[k]]) for k in list(BETA_COLUMNS.values())[1:])))
    return edition, out


# ------------------------------------------------------------------------------------------------ IMF DataMapper
def parse_datamapper(body: bytes, indicator: str, country: str = "IND") -> dict[int, Decimal]:
    """{year: value} for one indicator and country. Years before the edition's are actuals; the API doesn't say which
    years are forecasts."""
    import json
    try:
        js = json.loads(body)
    except ValueError as e:
        raise InputsError(f"not JSON ({str(e)[:60]})") from e
    values = (js.get("values") or {}).get(indicator) if isinstance(js, dict) else None
    series = (values or {}).get(country)
    if not isinstance(series, dict) or not series:
        raise InputsError(f"no values for {indicator} / {country}")
    out = {}
    for y, v in series.items():
        if not re.fullmatch(r"\d{4}", str(y)) or isinstance(v, bool) or not isinstance(v, (int, float)):
            raise InputsError(f"{indicator} / {country}: {y!r} -> {v!r} isn't a year and a number")
        out[int(y)] = Decimal(repr(float(v)))
    return dict(sorted(out.items()))


def load_weo_editions(path: Path | None = None) -> list[date]:
    """The WEO release dates in `weo_editions.yaml` (curated, each with its source), oldest first."""
    raw = yaml.safe_load((path or CONFIG_DIR / "weo_editions.yaml").read_text()) or {}
    out = []
    for e in raw.get("editions", []):
        d = _date(e.get("released")) if isinstance(e, dict) else None
        if d is None or not e.get("source"):
            raise InputsError(f"weo_editions.yaml: every edition needs a 'released' date and a 'source': {e!r}")
        out.append(d)
    return sorted(out)


@dataclass(frozen=True)
class Vintage:
    action: str                    # "none" (stored latest vintage stands), "write", "error"
    vintage: date | None
    note: str


def choose_vintage(editions: list[date], today: date, stored_latest: date | None, changed: bool) -> Vintage:
    """Which WEO edition the fetched values belong to. The API doesn't name it, so (decided by Claude, standing
    delegation, 28 Oct 2026): the newest configured edition released on or before today is the one the API should be
    serving. With nothing stored, the values are stored under that edition. With a stored vintage, unchanged values
    mean nothing to write (an API not yet updated shows up as a note); changed values are a new edition's, written
    under the newest configured edition that is later than the stored one, and if there is none (the YAML lacks the
    edition, or the API changed within an edition) the run says so and writes nothing."""
    current = [d for d in editions if d <= today]
    if not current:
        return Vintage("error", None, "no WEO edition in weo_editions.yaml is released yet")
    newest = current[-1]
    if stored_latest is None:
        return Vintage("write", newest, f"first load, stored as the {newest:%d %b %Y} edition")
    if not changed:
        if newest > stored_latest:
            return Vintage("none", None, f"the {newest:%d %b %Y} edition is out but the API still serves the "
                                         f"{stored_latest:%d %b %Y} values")
        return Vintage("none", None, "the stored edition's values are unchanged")
    later = [d for d in current if d > stored_latest]
    if not later:
        return Vintage("error", None, f"the values changed but no edition after {stored_latest:%d %b %Y} is in "
                                      "weo_editions.yaml (a new WEO needs its release date added there, with its source)")
    return Vintage("write", later[-1], f"new values, stored as the {later[-1]:%d %b %Y} edition")


# ------------------------------------------------------------------------------------------------ for 2A-13
@dataclass(frozen=True)
class Rate:
    value: Decimal                 # percent
    parts: dict                    # every input with its value and date, for the valuation run's `inputs`


INDIA_10Y = "fred.INDIRLTLT01STM"


def _latest(db, key: str, as_of: date) -> tuple[date, Decimal]:
    row = db.execute("""select o.obs_date, o.value from core.observation o join core.series s on s.id = o.series_id
                         where s.key = %s and o.obs_date <= %s order by o.obs_date desc limit 1""",
                     (key, as_of)).fetchone()
    if row is None:
        raise InputsError(f"no {key} observation on or before {as_of}")
    return row[0], row[1]


def inr_risk_free(db, as_of: date) -> Rate:
    """India's 10-year yield (FRED/OECD, monthly, the latest month on or before `as_of`) minus India's sovereign default
    spread from Damodaran's newest edition on or before `as_of` (Brief §6.3; Sep 2026: 7.06% - 1.75% = 5.31%)."""
    month, y10 = _latest(db, INDIA_10Y, as_of)
    edition, spread = _latest(db, "damodaran.IND_DEFAULT_SPREAD", as_of)
    return Rate(y10 - spread, {"india_10y_pct": y10, "india_10y_month": month.isoformat(), "default_spread_pct": spread,
                               "spread_edition": edition.isoformat()})


def india_erp(db, as_of: date) -> Rate:
    """The mature-market implied ERP of the latest month on or before `as_of` plus India's country risk premium from
    Damodaran's newest edition on or before `as_of` (never the file's total-ERP column; Oct 2026 mature 4.20% + 2.72%
    from the July edition = 6.92%)."""
    month, mature = _latest(db, "damodaran.ERP_IMPLIED", as_of)
    edition, crp = _latest(db, "damodaran.IND_CRP", as_of)
    return Rate(mature + crp, {"mature_erp_pct": mature, "mature_erp_month": month.isoformat(), "india_crp_pct": crp,
                               "crp_edition": edition.isoformat()})
